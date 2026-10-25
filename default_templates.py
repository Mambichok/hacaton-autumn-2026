import json
from dataclasses import dataclass
from datetime import datetime, timezone
from io import BytesIO
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt


@dataclass(frozen=True)
class TemplateField:
    name: str
    label: str
    type: str = "text"


@dataclass(frozen=True)
class DefaultTemplate:
    slug: str
    title: str
    description: str
    category: str
    fields: tuple[TemplateField, ...]
    recipient: str
    details: tuple[str, ...]


DEFAULT_TEMPLATES = (
    DefaultTemplate(
        "annual-leave",
        "Заявление на ежегодный оплачиваемый отпуск",
        "Типовая форма для запроса ежегодного оплачиваемого отпуска. Проверьте даты по графику отпусков.",
        "Кадровые",
        (
            TemplateField("recipient", "Кому (должность и ФИО)"),
            TemplateField("employee_name", "ФИО сотрудника"),
            TemplateField("position", "Должность"),
            TemplateField("department", "Подразделение"),
            TemplateField("leave_start", "Дата начала отпуска", "date"),
            TemplateField("leave_end", "Дата окончания отпуска", "date"),
            TemplateField("leave_days", "Количество календарных дней"),
            TemplateField("document_date", "Дата заявления", "date"),
        ),
        "руководителю организации",
        (
            "Прошу предоставить мне ежегодный оплачиваемый отпуск с {{leave_start}} по {{leave_end}} продолжительностью {{leave_days}} календарных дней.",
        ),
    ),
    DefaultTemplate(
        "unpaid-leave",
        "Заявление на отпуск без сохранения заработной платы",
        "Типовая форма отпуска без сохранения заработной платы по семейным или иным уважительным причинам.",
        "Кадровые",
        (
            TemplateField("recipient", "Кому (должность и ФИО)"),
            TemplateField("employee_name", "ФИО сотрудника"),
            TemplateField("position", "Должность"),
            TemplateField("department", "Подразделение"),
            TemplateField("leave_start", "Дата начала отпуска", "date"),
            TemplateField("leave_end", "Дата окончания отпуска", "date"),
            TemplateField("leave_days", "Количество календарных дней"),
            TemplateField("reason", "Причина отпуска"),
            TemplateField("document_date", "Дата заявления", "date"),
        ),
        "руководителю организации",
        (
            "Прошу предоставить мне отпуск без сохранения заработной платы с {{leave_start}} по {{leave_end}} продолжительностью {{leave_days}} календарных дней.",
            "Причина: {{reason}}.",
        ),
    ),
    DefaultTemplate(
        "resignation",
        "Заявление на увольнение по собственному желанию",
        "Типовая форма заявления работника о расторжении трудового договора по инициативе работника.",
        "Кадровые",
        (
            TemplateField("recipient", "Кому (должность и ФИО)"),
            TemplateField("employee_name", "ФИО сотрудника"),
            TemplateField("position", "Должность"),
            TemplateField("department", "Подразделение"),
            TemplateField("termination_date", "Предполагаемая дата увольнения", "date"),
            TemplateField("document_date", "Дата заявления", "date"),
        ),
        "руководителю организации",
        (
            "Прошу расторгнуть трудовой договор по моей инициативе и уволить меня {{termination_date}}.",
            "Прошу произвести окончательный расчет и выдать документы, связанные с работой, в установленном порядке.",
        ),
    ),
    DefaultTemplate(
        "position-transfer",
        "Заявление на перевод на другую должность или в другой отдел",
        "Типовая форма запроса на перевод; перевод оформляется при наличии согласия сторон и необходимых кадровых решений.",
        "Кадровые",
        (
            TemplateField("recipient", "Кому (должность и ФИО)"),
            TemplateField("employee_name", "ФИО сотрудника"),
            TemplateField("current_position", "Текущая должность"),
            TemplateField("current_department", "Текущее подразделение"),
            TemplateField("target_position", "Новая должность"),
            TemplateField("target_department", "Новое подразделение"),
            TemplateField("transfer_date", "Желаемая дата перевода", "date"),
            TemplateField("reason", "Причина / комментарий"),
            TemplateField("document_date", "Дата заявления", "date"),
        ),
        "руководителю организации",
        (
            "Прошу рассмотреть возможность моего перевода с должности {{current_position}} подразделения {{current_department}} на должность {{target_position}} подразделения {{target_department}} с {{transfer_date}}.",
            "Причина / комментарий: {{reason}}.",
        ),
    ),
    DefaultTemplate(
        "child-tax-deduction",
        "Заявление на стандартный налоговый вычет на детей",
        "Форма обращения за стандартным налоговым вычетом. Право и размер вычета подтвердите по актуальным правилам и документам.",
        "Кадровые",
        (
            TemplateField("recipient", "Кому (должность и ФИО)"),
            TemplateField("employee_name", "ФИО сотрудника"),
            TemplateField("position", "Должность"),
            TemplateField("department", "Подразделение"),
            TemplateField("children", "ФИО и даты рождения детей"),
            TemplateField("supporting_documents", "Прилагаемые подтверждающие документы"),
            TemplateField("document_date", "Дата заявления", "date"),
        ),
        "руководителю организации / в бухгалтерию",
        (
            "Прошу предоставить мне стандартный налоговый вычет по НДФЛ на детей в соответствии с действующим законодательством при наличии у меня права на такой вычет.",
            "Сведения о детях: {{children}}.",
            "Подтверждающие документы: {{supporting_documents}}.",
        ),
    ),
    DefaultTemplate(
        "advance-report",
        "Авансовый отчет",
        "Универсальная форма отчета о расходовании подотчетных средств. Состав приложений и учетные реквизиты проверьте по правилам организации.",
        "Финансы",
        (
            TemplateField("employee_name", "ФИО подотчетного лица"),
            TemplateField("position", "Должность"),
            TemplateField("department", "Подразделение"),
            TemplateField("advance_amount", "Получено под отчет (сумма)"),
            TemplateField("spent_amount", "Израсходовано (сумма)"),
            TemplateField("balance_amount", "Остаток к возврату (сумма)"),
            TemplateField("overrun_amount", "Перерасход к возмещению (сумма)"),
            TemplateField("expense_details", "Расшифровка расходов и реквизиты документов"),
            TemplateField("attachments", "Количество листов / приложенные документы"),
            TemplateField("document_date", "Дата отчета", "date"),
        ),
        "бухгалтерии",
        (
            "Отчет о расходовании денежных средств, выданных под отчет.",
            "Получено: {{advance_amount}}. Израсходовано: {{spent_amount}}. Остаток к возврату: {{balance_amount}}. Перерасход к возмещению: {{overrun_amount}}.",
            "Расшифровка расходов: {{expense_details}}.",
            "Приложено: {{attachments}}.",
        ),
    ),
    DefaultTemplate(
        "cash-advance-request",
        "Заявление на выдачу денежных средств под отчет",
        "Типовая заявка на выдачу средств для согласованных хозяйственных или служебных расходов.",
        "Финансы",
        (
            TemplateField("recipient", "Кому (должность и ФИО)"),
            TemplateField("employee_name", "ФИО сотрудника"),
            TemplateField("position", "Должность"),
            TemplateField("department", "Подразделение"),
            TemplateField("amount", "Запрашиваемая сумма"),
            TemplateField("purpose", "Цель расходов"),
            TemplateField("issue_date", "Желаемая дата выдачи", "date"),
            TemplateField("report_due_date", "Срок представления отчета", "date"),
            TemplateField("document_date", "Дата заявления", "date"),
        ),
        "руководителю организации / в бухгалтерию",
        (
            "Прошу выдать мне под отчет денежные средства в размере {{amount}} для следующей цели: {{purpose}}.",
            "Желаемая дата выдачи: {{issue_date}}. Срок представления подтверждающих документов и отчета: {{report_due_date}}.",
        ),
    ),
    DefaultTemplate(
        "banking-access",
        "Заявка на предоставление или расширение прав в банковских системах",
        "Внутренняя заявка на банковскую систему; конкретные роли и согласования определяются владельцем системы и политиками банка.",
        "ИТ и доступы",
        (
            TemplateField("employee_name", "ФИО пользователя"),
            TemplateField("position", "Должность"),
            TemplateField("department", "Подразделение"),
            TemplateField("system_name", "Банковская система"),
            TemplateField("organization_account", "Организация / учетная запись"),
            TemplateField("access_action", "Действие: создать / изменить / закрыть"),
            TemplateField("requested_role", "Запрашиваемая роль и полномочия"),
            TemplateField("reason", "Обоснование служебной необходимости"),
            TemplateField("access_start", "Дата начала доступа", "date"),
            TemplateField("access_end", "Дата окончания доступа", "date"),
            TemplateField("approver", "Руководитель / согласующий"),
            TemplateField("document_date", "Дата заявки", "date"),
        ),
        "владельцу банковской системы / подразделению сопровождения",
        (
            "Прошу {{access_action}} права доступа пользователю {{employee_name}} ({{position}}, {{department}}) в системе {{system_name}} для организации / учетной записи {{organization_account}}.",
            "Необходимая роль и полномочия: {{requested_role}}.",
            "Служебное обоснование: {{reason}}.",
            "Период доступа: с {{access_start}} по {{access_end}}. Согласующий руководитель: {{approver}}.",
        ),
    ),
    DefaultTemplate(
        "vpn-access",
        "Служебная записка на удаленный доступ (VPN)",
        "Внутренняя заявка на удаленный доступ с указанием периода и рабочей необходимости.",
        "ИТ и доступы",
        (
            TemplateField("recipient", "Кому (должность / подразделение)"),
            TemplateField("employee_name", "ФИО сотрудника"),
            TemplateField("position", "Должность"),
            TemplateField("department", "Подразделение"),
            TemplateField("access_start", "Начало периода доступа", "date"),
            TemplateField("access_end", "Окончание периода доступа", "date"),
            TemplateField("resources", "Необходимые ресурсы / системы"),
            TemplateField("reason", "Рабочая необходимость"),
            TemplateField("manager_name", "ФИО согласующего руководителя"),
            TemplateField("document_date", "Дата записки", "date"),
        ),
        "подразделению информационных технологий",
        (
            "Прошу предоставить удаленный доступ (VPN) сотруднику {{employee_name}}, {{position}}, {{department}}, на период с {{access_start}} по {{access_end}}.",
            "Необходимые рабочие ресурсы: {{resources}}.",
            "Обоснование: {{reason}}.",
            "Согласующий руководитель: {{manager_name}}.",
        ),
    ),
    DefaultTemplate(
        "account-request",
        "Заявка на создание или изменение учетной записи",
        "Внутренняя заявка для создания, изменения или блокировки учетной записи и назначения рабочих прав.",
        "ИТ и доступы",
        (
            TemplateField("recipient", "Кому (подразделение / ответственное лицо)"),
            TemplateField("employee_name", "ФИО пользователя"),
            TemplateField("position", "Должность"),
            TemplateField("department", "Подразделение"),
            TemplateField("system_name", "Система / сервис"),
            TemplateField("account_action", "Действие: создать / изменить / заблокировать"),
            TemplateField("login", "Логин (если известен)"),
            TemplateField("requested_role", "Роль / необходимые права"),
            TemplateField("effective_date", "Дата выполнения", "date"),
            TemplateField("reason", "Основание / комментарий"),
            TemplateField("manager_name", "ФИО согласующего руководителя"),
            TemplateField("document_date", "Дата заявки", "date"),
        ),
        "подразделению информационных технологий",
        (
            "Прошу {{account_action}} учетную запись пользователя {{employee_name}}, {{position}}, {{department}}, в системе {{system_name}}.",
            "Логин: {{login}}. Требуемая роль / права: {{requested_role}}.",
            "Основание: {{reason}}. Согласующий руководитель: {{manager_name}}.",
            "Желаемая дата выполнения: {{effective_date}}.",
        ),
    ),
)


def _build_document(template: DefaultTemplate) -> bytes:
    document = Document()
    section = document.sections[0]
    section.top_margin = Inches(0.75)
    section.bottom_margin = Inches(0.75)
    section.left_margin = Inches(1)
    section.right_margin = Inches(0.75)
    normal = document.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(12)

    field_names = {field.name for field in template.fields}
    if "recipient" in field_names:
        document.add_paragraph("Кому: {{recipient}}").alignment = WD_ALIGN_PARAGRAPH.RIGHT
    document.add_paragraph("От: {{employee_name}}").alignment = WD_ALIGN_PARAGRAPH.RIGHT
    if "position" in field_names:
        document.add_paragraph("Должность: {{position}}").alignment = WD_ALIGN_PARAGRAPH.RIGHT
    elif "current_position" in field_names:
        document.add_paragraph(
            "Текущая должность: {{current_position}}"
        ).alignment = WD_ALIGN_PARAGRAPH.RIGHT
    if "department" in field_names:
        document.add_paragraph("Подразделение: {{department}}").alignment = WD_ALIGN_PARAGRAPH.RIGHT
    elif "current_department" in field_names:
        document.add_paragraph(
            "Текущее подразделение: {{current_department}}"
        ).alignment = WD_ALIGN_PARAGRAPH.RIGHT
    heading = document.add_paragraph()
    heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = heading.add_run(template.title.upper())
    run.bold = True
    run.font.size = Pt(14)

    for detail in template.details:
        document.add_paragraph(detail)

    document.add_paragraph()
    document.add_paragraph("Дата: {{document_date}}")
    document.add_paragraph("Подпись: ____________________ / {{employee_name}} /")
    document.add_paragraph("Согласование: ____________________")
    output = BytesIO()
    document.save(output)
    return output.getvalue()


def seed_default_templates(connection, template_directory: Path) -> int:
    if not connection.in_transaction:
        connection.execute("BEGIN IMMEDIATE")
    connection.execute(
        """CREATE TABLE IF NOT EXISTS app_settings (
            setting_key TEXT PRIMARY KEY,
            setting_value TEXT NOT NULL
        )"""
    )
    seeded = connection.execute(
        "SELECT setting_value FROM app_settings WHERE setting_key = ?",
        ("default_templates_version",),
    ).fetchone()
    if seeded:
        return 0

    template_directory.mkdir(parents=True, exist_ok=True)
    created_at = datetime.now(timezone.utc).isoformat()
    for template in DEFAULT_TEMPLATES:
        template_id = f"builtin-{template.slug}"
        path = template_directory / f"{template_id}.docx"
        path.write_bytes(_build_document(template))
        fields = [
            {
                "name": field.name,
                "label": (
                    f"Кому ({template.recipient})"
                    if field.name == "recipient"
                    else field.label
                ),
                "type": field.type,
            }
            for field in template.fields
        ]
        connection.execute(
            "INSERT INTO templates "
            "(template_id, title, description, category, fields_json, created_at) "
            "VALUES (?, ?, ?, ?, ?, ?)",
            (
                template_id,
                template.title,
                template.description,
                template.category,
                json.dumps(fields, ensure_ascii=False),
                created_at,
            ),
        )
    connection.execute(
        "INSERT INTO app_settings (setting_key, setting_value) VALUES (?, ?)",
        ("default_templates_version", "1"),
    )
    return len(DEFAULT_TEMPLATES)
