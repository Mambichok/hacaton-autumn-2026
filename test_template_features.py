import json
import sqlite3
import tempfile
import unittest
from types import SimpleNamespace
from io import BytesIO
from pathlib import Path
from unittest.mock import patch

from docx import Document
from langchain_core.documents import Document as SearchDocument
from fastapi.testclient import TestClient

import app as application
import default_templates
from ask_bot import BotServiceError, ask_bot, review_draft
import index_documents


class TemplateFeatureTests(unittest.TestCase):
    def setUp(self):
        self.temporary_directory = tempfile.TemporaryDirectory()
        self.addCleanup(self.temporary_directory.cleanup)
        self.data_directory = Path(self.temporary_directory.name)
        self.client = TestClient(application.app)
        self.patches = (
            patch.object(application, "TEMPLATE_DIR", self.data_directory / "templates"),
            patch.object(application, "ADMIN_DATABASE", self.data_directory / "admins.sqlite3"),
            patch.object(application, "_current_user", return_value={"username": "admin"}),
        )
        for patcher in self.patches:
            patcher.start()
            self.addCleanup(patcher.stop)

    @staticmethod
    def make_template() -> bytes:
        document = Document()
        paragraph = document.add_paragraph()
        paragraph.add_run("Сотрудник: {{empl")
        paragraph.add_run("oyee_name}}")
        document.add_paragraph("Дата: {{document_date}}")
        output = BytesIO()
        document.save(output)
        return output.getvalue()

    def test_uploaded_template_generates_completed_docx(self):
        fields = [
            {"name": "employee_name", "label": "ФИО сотрудника", "type": "text"},
            {"name": "document_date", "label": "Дата документа", "type": "date"},
        ]
        response = self.client.post(
            "/api/templates",
            data={
                "title": "Служебная записка",
                "description": "Внутренний документ",
                "fields": json.dumps(fields, ensure_ascii=False),
            },
            files={
                "file": (
                    "memo.docx",
                    self.make_template(),
                    "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                )
            },
        )
        self.assertEqual(response.status_code, 201, response.text)
        template = response.json()["template"]
        self.assertEqual(template["category"], "Общие")

        generated = self.client.post(
            f"/api/templates/{template['id']}/generate",
            json={"values": {"employee_name": "Иванов Иван", "document_date": "2025-06-07"}},
        )

        self.assertEqual(generated.status_code, 200, generated.text)
        self.assertIn("application/vnd.openxmlformats-officedocument.wordprocessingml.document", generated.headers["content-type"])
        self.assertIn("attachment", generated.headers["content-disposition"])
        self.assertTrue(generated.headers["content-disposition"].endswith(".docx"))
        document = Document(BytesIO(generated.content))
        text = "\n".join(paragraph.text for paragraph in document.paragraphs)
        self.assertIn("Сотрудник: Иванов Иван", text)
        self.assertIn("Дата: 07.06.2025", text)
        self.assertNotIn("{{", text)

    def test_template_upload_rejects_fields_not_present_in_docx(self):
        response = self.client.post(
            "/api/templates",
            data={
                "title": "Служебная записка",
                "description": "",
                "fields": json.dumps(
                    [{"name": "unknown", "label": "Неизвестно", "type": "text"}]
                ),
            },
            files={"file": ("memo.docx", self.make_template())},
        )

        self.assertEqual(response.status_code, 400)
        self.assertIn("Не совпадают поля", response.json()["detail"])

    def test_admin_can_inspect_docx_and_create_fields_from_selected_text(self):
        document = Document()
        paragraph = document.add_paragraph()
        paragraph.add_run("Сотрудник: ")
        paragraph.add_run("Иванов")
        paragraph.add_run(" Иван Иванович")
        output = BytesIO()
        document.save(output)
        source = output.getvalue()

        inspected = self.client.post(
            "/api/templates/inspect",
            files={"file": ("memo.docx", source)},
        )

        self.assertEqual(inspected.status_code, 200, inspected.text)
        self.assertEqual(inspected.json()["paragraphs"][0]["index"], 0)
        self.assertEqual(
            inspected.json()["paragraphs"][0]["text"],
            "Сотрудник: Иванов Иван Иванович",
        )

        start = len("Сотрудник: ")
        uploaded = self.client.post(
            "/api/templates",
            data={
                "title": "Служебная записка",
                "description": "",
                "fields": json.dumps(
                    [{
                        "name": "field_1",
                        "label": "ФИО сотрудника",
                        "type": "text",
                        "paragraph_index": 0,
                        "start": start,
                        "end": start + len("Иванов Иван Иванович"),
                    }],
                    ensure_ascii=False,
                ),
            },
            files={"file": ("memo.docx", source)},
        )

        self.assertEqual(uploaded.status_code, 201, uploaded.text)
        template_id = uploaded.json()["template"]["id"]
        generated = self.client.post(
            f"/api/templates/{template_id}/generate",
            json={"values": {"field_1": "Петрова Анна"}},
        )
        self.assertEqual(generated.status_code, 200, generated.text)
        generated_document = Document(BytesIO(generated.content))
        self.assertEqual(
            generated_document.paragraphs[0].text,
            "Сотрудник: Петрова Анна",
        )

    def test_default_templates_are_seeded_once_and_generate_docx(self):
        with application._admin_database() as connection:
            seeded_count = default_templates.seed_default_templates(
                connection,
                self.data_directory / "templates",
            )
        self.assertEqual(seeded_count, 10)

        templates = self.client.get("/api/templates").json()["templates"]
        self.assertEqual(len(templates), 10)
        self.assertEqual(
            {template["title"] for template in templates},
            {template.title for template in default_templates.DEFAULT_TEMPLATES},
        )

        for template in templates:
            values = {
                field["name"]: "2025-06-07" if field["type"] == "date" else "Тестовое значение"
                for field in template["fields"]
            }
            response = self.client.post(
                f"/api/templates/{template['id']}/generate",
                json={"values": values},
            )
            self.assertEqual(response.status_code, 200, f"{template['title']}: {response.text}")
            generated_document = Document(BytesIO(response.content))
            document_text = "\n".join(
                paragraph.text
                for paragraph in application._iter_docx_paragraphs(generated_document)
            )
            self.assertNotIn("{{", document_text, template["title"])
            self.assertIn("Тестовое значение", document_text, template["title"])

        with application._admin_database() as connection:
            deleted_id = "builtin-annual-leave"
            connection.execute("DELETE FROM templates WHERE template_id = ?", (deleted_id,))
            seeded_again = default_templates.seed_default_templates(
                connection,
                self.data_directory / "templates",
            )
        self.assertEqual(seeded_again, 0)

    def test_application_startup_installs_default_templates(self):
        with TestClient(application.app) as startup_client:
            response = startup_client.get("/api/templates")

        self.assertEqual(response.status_code, 200)
        self.assertEqual(len(response.json()["templates"]), 10)
        self.assertTrue(
            (self.data_directory / "templates" / "builtin-annual-leave.docx").is_file()
        )

    def test_template_generation_requires_every_field(self):
        fields = [
            {"name": "employee_name", "label": "ФИО сотрудника", "type": "text"},
            {"name": "document_date", "label": "Дата документа", "type": "date"},
        ]
        uploaded = self.client.post(
            "/api/templates",
            data={
                "title": "Служебная записка",
                "description": "",
                "fields": json.dumps(fields, ensure_ascii=False),
            },
            files={"file": ("memo.docx", self.make_template())},
        )
        template_id = uploaded.json()["template"]["id"]

        response = self.client.post(
            f"/api/templates/{template_id}/generate",
            json={"values": {"employee_name": "Иванов Иван"}},
        )

        self.assertEqual(response.status_code, 400)

    def test_admin_can_delete_uploaded_template(self):
        fields = [
            {"name": "employee_name", "label": "ФИО сотрудника", "type": "text"},
            {"name": "document_date", "label": "Дата документа", "type": "date"},
        ]
        uploaded = self.client.post(
            "/api/templates",
            data={
                "title": "Служебная записка",
                "description": "",
                "fields": json.dumps(fields, ensure_ascii=False),
            },
            files={"file": ("memo.docx", self.make_template())},
        )
        self.assertEqual(uploaded.status_code, 201, uploaded.text)
        template = uploaded.json()["template"]
        template_path = self.data_directory / "templates" / f"{template['id']}.docx"
        self.assertTrue(template_path.is_file())

        deleted = self.client.delete(f"/api/templates/{template['id']}")
        listed = self.client.get("/api/templates")

        self.assertEqual(deleted.status_code, 204)
        self.assertFalse(template_path.exists())
        self.assertEqual(listed.json()["templates"], [])

    def test_existing_template_database_is_migrated_to_categories(self):
        legacy_database = self.data_directory / "legacy.sqlite3"
        connection = sqlite3.connect(legacy_database)
        try:
            connection.execute(
                """CREATE TABLE templates (
                    template_id TEXT PRIMARY KEY,
                    title TEXT NOT NULL,
                    description TEXT NOT NULL,
                    fields_json TEXT NOT NULL,
                    created_at TEXT NOT NULL
                )"""
            )
            connection.execute(
                "INSERT INTO templates VALUES (?, ?, ?, ?, ?)",
                ("legacy", "Old template", "", "[]", "2026-01-01"),
            )
            connection.commit()
        finally:
            connection.close()

        with patch.object(application, "ADMIN_DATABASE", legacy_database):
            with application._admin_database() as connection:
                row = connection.execute(
                    "SELECT category FROM templates WHERE template_id = 'legacy'"
                ).fetchone()

        self.assertEqual(row["category"], "Общие")

    def test_document_read_route_uses_inline_content_disposition(self):
        original_directory = self.data_directory / "original"
        original_directory.mkdir()
        (original_directory / "policy.docx").write_bytes(b"docx")
        with (
            patch.object(application, "ORIGINAL_DIR", original_directory),
            patch.object(application, "UPLOAD_DIR", self.data_directory / "uploaded"),
        ):
            response = self.client.get("/api/documents/policy.docx/view")

        self.assertEqual(response.status_code, 200)
        self.assertIn("inline", response.headers["content-disposition"])
        self.assertIn(
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            response.headers["content-type"],
        )

    def test_empty_document_list_explains_missing_source_files_and_is_not_cached(self):
        with (
            patch.object(application, "ORIGINAL_DIR", self.data_directory / "missing-original"),
            patch.object(application, "UPLOAD_DIR", self.data_directory / "missing-upload"),
        ):
            response = self.client.get("/api/documents")

        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.json()["documents"], [])
        self.assertIn("DOCUMENTS_ORIGINAL_DIR", response.json()["detail"])
        self.assertEqual(response.headers["cache-control"], "no-store")

    def test_document_check_defaults_to_reviewing_against_the_database(self):
        with (
            patch.object(application, "read_docx_bytes", return_value="Текст проекта."),
            patch.object(
                application,
                "review_draft",
                return_value={"answer": "Проверено по базе.", "sources": []},
            ) as review,
        ):
            response = self.client.post(
                "/api/documents/check",
                files={"file": ("draft.docx", b"docx")},
                data={"mode": "spelling"},
            )

        self.assertEqual(response.status_code, 200, response.text)
        self.assertEqual(response.json()["answer"], "Проверено по базе.")
        review.assert_called_once_with("Текст проекта.")

    def test_chat_error_explains_search_dependency(self):
        detail = application._chat_error_detail(
            BotServiceError("search", RuntimeError("connection refused"))
        )

        self.assertIn("Qdrant", detail)
        self.assertIn("эмбеддингов", detail)

    def test_chat_error_explains_rejected_openrouter_key(self):
        cause = RuntimeError("unauthorized")
        cause.status_code = 401
        detail = application._chat_error_detail(BotServiceError("model", cause))

        self.assertIn("OPENROUTER_API_KEY", detail)

    def test_reindex_uses_new_collection_before_switching_active_alias(self):
        class FakeClient:
            aliases = [SimpleNamespace(
                alias_name=index_documents.ACTIVE_COLLECTION_NAME,
                collection_name="documents_build_previous",
            )]

            def get_aliases(self):
                return SimpleNamespace(aliases=self.aliases)

            def count(self, collection_name, exact):
                self.counted_collection = collection_name
                self.exact = exact
                return SimpleNamespace(count=2)

            def update_collection_aliases(self, operations):
                self.operations = operations

            def collection_exists(self, collection_name):
                return True

            def delete_collection(self, collection_name):
                self.deleted_collection = collection_name

        client = FakeClient()
        with patch.object(index_documents.QdrantVectorStore, "from_texts") as build:
            index_documents.replace_active_collection(
                client,
                ["part one", "part two"],
                [{"chunk": 0}, {"chunk": 1}],
                object(),
            )

        staged_name = build.call_args.kwargs["collection_name"]
        self.assertTrue(staged_name.startswith("documents_build_"))
        self.assertEqual(client.counted_collection, staged_name)
        self.assertTrue(client.exact)
        self.assertEqual(len(client.operations), 2)
        self.assertEqual(
            client.operations[-1].create_alias.alias_name,
            index_documents.ACTIVE_COLLECTION_NAME,
        )
        self.assertEqual(client.operations[-1].create_alias.collection_name, staged_name)
        self.assertEqual(client.deleted_collection, "documents_build_previous")

    def test_failed_reindex_keeps_existing_active_collection(self):
        class FakeClient:
            def __init__(self):
                self.aliases = [SimpleNamespace(
                    alias_name=index_documents.ACTIVE_COLLECTION_NAME,
                    collection_name="documents_build_previous",
                )]
                self.alias_update_called = False
                self.deleted_collections = []

            def get_aliases(self):
                return SimpleNamespace(aliases=self.aliases)

            def collection_exists(self, collection_name):
                return True

            def update_collection_aliases(self, operations):
                self.alias_update_called = True

            def delete_collection(self, collection_name):
                self.deleted_collections.append(collection_name)

        client = FakeClient()
        with (
            patch.object(
                index_documents.QdrantVectorStore,
                "from_texts",
                side_effect=RuntimeError("embedding provider is unavailable"),
            ),
            self.assertRaisesRegex(RuntimeError, "embedding provider is unavailable"),
        ):
            index_documents.replace_active_collection(
                client,
                ["part one"],
                [{"chunk": 0}],
                object(),
            )

        self.assertFalse(client.alias_update_called)
        self.assertEqual(
            client.aliases[0].collection_name,
            "documents_build_previous",
        )

    def test_reindex_includes_original_and_uploaded_docx_and_pdf(self):
        originals = self.data_directory / "original"
        uploads = self.data_directory / "uploaded"
        originals.mkdir()
        uploads.mkdir()
        original_docx = originals / "policy.docx"
        uploaded_pdf = uploads / "notice.pdf"
        uploaded_docx = uploads / "memo.docx"
        ignored = uploads / "notes.txt"
        for file in (original_docx, uploaded_pdf, uploaded_docx, ignored):
            file.write_bytes(b"test")

        with (
            patch.object(index_documents, "DOCUMENTS_DIR", originals),
            patch.object(index_documents, "UPLOAD_DIR", uploads),
        ):
            files = index_documents.get_document_files()

        self.assertEqual(
            {file.name for file in files},
            {"policy.docx", "notice.pdf", "memo.docx"},
        )

    def test_chat_uses_longer_history_without_embedding_the_entire_history(self):
        source = SearchDocument(
            page_content="Сотрудники подают документы по установленной форме.",
            metadata={"source": "procedure.docx", "reference": "Пункт 1"},
        )

        class SearchStore:
            query = ""

            def similarity_search_with_relevance_scores(self, query, k):
                self.query = query
                return [(source, 0.9)]

        store = SearchStore()
        history = [
            {"role": "user", "content": f"История диалога {index}"}
            for index in range(12)
        ]
        with (
            patch("ask_bot.get_vector_store", return_value=store),
            patch("ask_bot._call_model", return_value="Ответ[[SOURCES:S1]]") as call_model,
        ):
            result = ask_bot("Как подать документы?", history)

        self.assertLess(len(store.query), 4000)
        self.assertIn("История диалога 11", call_model.call_args.args[1])
        self.assertEqual(result["answer"], "Ответ")
        self.assertEqual(result["sources"][0]["title"], "procedure.docx")

    def test_document_review_searches_several_document_sections(self):
        class SearchStore:
            queries = []

            def similarity_search_with_relevance_scores(self, query, k):
                self.queries.append(query)
                return [
                    (
                        SearchDocument(
                            page_content=f"Требование из раздела {len(self.queries)}.",
                            metadata={
                                "source": f"policy-{len(self.queries)}.docx",
                                "reference": "Пункт 1",
                                "chunk": len(self.queries),
                            },
                        ),
                        0.9,
                    )
                ]

        store = SearchStore()
        draft = "Первый раздел.\n" + ("Содержимое документа. " * 400) + "\nПоследний раздел."
        with (
            patch("ask_bot.get_vector_store", return_value=store),
            patch("ask_bot._call_model", return_value="Найдены замечания.") as call_model,
        ):
            result = review_draft(draft)

        self.assertEqual(len(store.queries), 3)
        self.assertIn(draft[:100], store.queries[0])
        self.assertIn(draft[-100:], store.queries[-1])
        self.assertIn("Найдены замечания", result["answer"])
        self.assertEqual(len(result["sources"]), 3)
        self.assertIn("Сопоставь стиль", call_model.call_args.args[1])


if __name__ == "__main__":
    unittest.main()
