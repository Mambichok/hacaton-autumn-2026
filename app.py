import base64
import hashlib
import hmac
import json
import logging
import math
import os
import re
import secrets
import sqlite3
import threading
import time
import unicodedata
from contextlib import contextmanager
from datetime import datetime, timezone
from io import BytesIO
from pathlib import Path
from typing import Literal
from urllib.parse import quote
import uuid

from dotenv import load_dotenv
from fastapi import FastAPI, File, Form, HTTPException, Request, UploadFile
from fastapi.responses import FileResponse, JSONResponse, Response, StreamingResponse
from docx import Document
from docx.table import Table
from langchain_qdrant import QdrantVectorStore
from pydantic import BaseModel, Field
from qdrant_client import QdrantClient

from ask_bot import ask_bot, review_draft
from default_templates import seed_default_templates
from index_documents import (
    ACTIVE_COLLECTION_NAME,
    QDRANT_URL,
    create_splitter,
    extract_reference,
    extract_version,
    get_embeddings,
    ensure_active_collection,
    read_docx_bytes,
    read_pdf_bytes,
    reindex_documents,
)

load_dotenv()
logging.basicConfig(level=logging.INFO)
logger = logging.getLogger(__name__)

APP_DIR = Path(__file__).resolve().parent
ORIGINAL_DIR = Path(
    os.getenv("DOCUMENTS_ORIGINAL_DIR", str(APP_DIR / "documents" / "original"))
)
UPLOAD_DIR = Path(
    os.getenv("DOCUMENTS_UPLOAD_DIR", str(APP_DIR / "documents" / "uploaded"))
)
TEMPLATE_DIR = Path(
    os.getenv("DOCUMENT_TEMPLATES_DIR", str(APP_DIR / "documents" / "templates"))
)
ADMIN_DATABASE = Path(
    os.getenv("ADMIN_DATABASE_PATH", str(APP_DIR / "documents" / "admins.sqlite3"))
)
MAX_UPLOAD_SIZE = int(os.getenv("MAX_UPLOAD_SIZE_MB", "20")) * 1024 * 1024
ALLOWED_EXTENSIONS = {".docx", ".pdf"}
DOCUMENT_INDEX_LOCK = threading.Lock()
SESSION_COOKIE = "assistant_session"
SESSION_DURATION_SECONDS = 8 * 60 * 60
SENSITIVE_DATA_PATTERNS = (
    ("email", re.compile(r"\b[A-Z0-9._%+-]+@[A-Z0-9.-]+\.[A-Z]{2,}\b", re.IGNORECASE)),
    (
        "телефон",
        re.compile(
            r"(?<!\d)(?:\+7|8)[\s(.-]*\d{3}[\s).-]*\d{3}[\s.-]*\d{2}[\s.-]*\d{2}(?!\d)"
        ),
    ),
    ("СНИЛС", re.compile(r"(?<!\d)\d{3}[- ]?\d{3}[- ]?\d{3}[- ]?\d{2}(?!\d)")),
    ("ИНН", re.compile(r"(?<!\d)\d{10}(?:\d{2})?(?!\d)")),
    (
        "паспортные данные",
        re.compile(r"паспорт.{0,20}\b\d{2}\s?\d{2}\s?\d{6}\b", re.IGNORECASE),
    ),
)

app = FastAPI(title="Document Assistant")


def _seed_default_templates() -> None:
    with _admin_database() as connection:
        created = seed_default_templates(connection, TEMPLATE_DIR)
    if created:
        logger.info("Installed %s default DOCX templates", created)


app.add_event_handler("startup", _seed_default_templates)


class InMemoryRateLimiter:
    def __init__(self):
        self._buckets: dict[tuple[str, str], tuple[float, int]] = {}
        self._lock = threading.Lock()

    def check(self, scope: str, identity: str, limit: int, window_seconds: int) -> int:
        now = time.monotonic()
        key = (scope, identity)
        with self._lock:
            expired = [
                bucket_key
                for bucket_key, (started_at, _) in self._buckets.items()
                if now - started_at >= window_seconds
            ]
            for bucket_key in expired:
                del self._buckets[bucket_key]

            started_at, count = self._buckets.get(key, (now, 0))
            if count >= limit:
                return max(1, math.ceil(window_seconds - (now - started_at)))
            if len(self._buckets) >= 10_000 and key not in self._buckets:
                oldest_key = min(self._buckets, key=lambda item: self._buckets[item][0])
                del self._buckets[oldest_key]
            self._buckets[key] = (started_at, count + 1)
            return 0


rate_limiter = InMemoryRateLimiter()


def _request_identity(request: Request) -> str:
    return request.client.host if request.client else "unknown"


def _enforce_rate_limit(
    request: Request, scope: str, limit: int, window_seconds: int
) -> None:
    retry_after = rate_limiter.check(
        scope, _request_identity(request), limit, window_seconds
    )
    if retry_after:
        raise HTTPException(
            status_code=429,
            detail="Слишком много запросов. Подождите и попробуйте снова.",
            headers={"Retry-After": str(retry_after)},
        )


def _find_sensitive_data(text: str) -> list[str]:
    return [name for name, pattern in SENSITIVE_DATA_PATTERNS if pattern.search(text)]


def _embedding_quota_error_detail(exc: Exception) -> str | None:
    if "free-models-per-day" not in str(exc):
        return None
    return (
        "Исчерпан дневной лимит бесплатной модели эмбеддингов OpenRouter. "
        "Повторите попытку после сброса лимита или настройте модель с доступной квотой. "
        "При смене модели эмбеддингов потребуется переиндексация документов."
    )


def _chat_error_detail(exc: Exception) -> str:
    cause = getattr(exc, "cause", exc)
    quota_detail = _embedding_quota_error_detail(cause)
    if quota_detail:
        return quota_detail

    status_code = getattr(cause, "status_code", None)
    stage = getattr(exc, "stage", "")
    if status_code == 401:
        return "OpenRouter отклонил API-ключ. Проверьте OPENROUTER_API_KEY в .env."
    if status_code == 402:
        return "OpenRouter сообщает о недостатке средств или доступных кредитов."
    if status_code == 404:
        return "Выбранная модель не найдена или недоступна. Проверьте OPENAI_MODEL в .env."
    if status_code == 429:
        return "OpenRouter ограничил запрос по квоте или частоте. Проверьте лимиты модели и повторите позже."
    if isinstance(status_code, int) and status_code >= 500:
        return "Сервис OpenRouter временно недоступен. Повторите запрос позже."
    if stage == "search":
        return (
            "Не удалось найти контекст в базе документов. Проверьте, что Qdrant запущен, "
            "коллекция documents существует, а ключ и модель эмбеддингов доступны."
        )
    if stage == "model":
        return (
            "Не удалось получить ответ от модели OpenRouter. Проверьте доступность API-ключа "
            "и значение OPENAI_MODEL в .env."
        )
    return "Не удалось получить ответ. Проверьте журнал сервера и доступность Qdrant и OpenRouter."


class ChatMessage(BaseModel):
    role: Literal["user", "assistant"]
    content: str = Field(max_length=1200)


class ChatRequest(BaseModel):
    question: str = Field(min_length=1, max_length=4000)
    history: list[ChatMessage] = Field(default_factory=list, max_length=12)


class TemplateGenerateRequest(BaseModel):
    values: dict[str, str]


class LoginRequest(BaseModel):
    username: str = Field(min_length=1, max_length=80)
    password: str = Field(min_length=1, max_length=256)


class AdminCreateRequest(BaseModel):
    username: str = Field(min_length=3, max_length=80, pattern=r"^[A-Za-z0-9._@-]+$")
    password: str = Field(min_length=14, max_length=256)


def _connect_admin_database() -> sqlite3.Connection:
    ADMIN_DATABASE.parent.mkdir(parents=True, exist_ok=True)
    connection = sqlite3.connect(ADMIN_DATABASE, timeout=10)
    connection.row_factory = sqlite3.Row
    connection.execute(
        """CREATE TABLE IF NOT EXISTS admins (
            username TEXT PRIMARY KEY COLLATE NOCASE,
            password_hash TEXT NOT NULL,
            created_at TEXT NOT NULL,
            created_by TEXT NOT NULL
        )"""
    )
    connection.execute(
        """CREATE TABLE IF NOT EXISTS templates (
            template_id TEXT PRIMARY KEY,
            title TEXT NOT NULL,
            description TEXT NOT NULL,
            category TEXT NOT NULL DEFAULT 'Общие',
            fields_json TEXT NOT NULL,
            created_at TEXT NOT NULL
        )"""
    )
    connection.execute(
        """CREATE TABLE IF NOT EXISTS document_metadata (
            source TEXT PRIMARY KEY,
            title TEXT NOT NULL,
            updated_at TEXT NOT NULL
        )"""
    )
    template_columns = {
        column["name"]
        for column in connection.execute("PRAGMA table_info(templates)")
    }
    if "category" not in template_columns:
        connection.execute(
            "ALTER TABLE templates ADD COLUMN category TEXT NOT NULL DEFAULT 'Общие'"
        )
    return connection


@contextmanager
def _admin_database():
    connection = _connect_admin_database()
    try:
        yield connection
        connection.commit()
    except Exception:
        connection.rollback()
        raise
    finally:
        connection.close()


def _session_secret() -> bytes:
    secret = os.getenv("ADMIN_SESSION_SECRET", "")
    if len(secret) < 32:
        raise HTTPException(
            status_code=503,
            detail="Задайте ADMIN_SESSION_SECRET длиной не менее 32 символов в .env.",
        )
    return secret.encode("utf-8")


def _admin_password_hash(password: str, salt: bytes | None = None) -> str:
    salt = salt or secrets.token_bytes(16)
    digest = hashlib.scrypt(
        password.encode("utf-8"),
        salt=salt,
        n=2**15,
        r=8,
        p=1,
        dklen=32,
        maxmem=64 * 1024 * 1024,
    )
    return f"{salt.hex()}:{digest.hex()}"


def _verify_admin_password(password: str, stored_hash: str) -> bool:
    try:
        salt_hex, digest_hex = stored_hash.split(":", 1)
        expected = bytes.fromhex(digest_hex)
        actual_hash = _admin_password_hash(password, bytes.fromhex(salt_hex))
        actual = bytes.fromhex(actual_hash.split(":", 1)[1])
    except (ValueError, TypeError):
        return False
    return hmac.compare_digest(actual, expected)


def _issue_session(username: str, role: str) -> str:
    payload = {
        "sub": username,
        "role": role,
        "exp": int(datetime.now(timezone.utc).timestamp()) + SESSION_DURATION_SECONDS,
    }
    encoded = base64.urlsafe_b64encode(
        json.dumps(payload, separators=(",", ":")).encode("utf-8")
    ).rstrip(b"=")
    signature = hmac.new(_session_secret(), encoded, hashlib.sha256).digest()
    return (encoded + b"." + base64.urlsafe_b64encode(signature).rstrip(b"=")).decode("ascii")


def _current_user(request: Request) -> dict[str, str]:
    token = request.cookies.get(SESSION_COOKIE)
    if not token:
        raise HTTPException(status_code=401, detail="Требуется вход администратора.")
    try:
        encoded, signature = token.encode("ascii").split(b".", 1)
        expected = hmac.new(_session_secret(), encoded, hashlib.sha256).digest()
        supplied = base64.urlsafe_b64decode(signature + b"=" * (-len(signature) % 4))
        if not hmac.compare_digest(expected, supplied):
            raise ValueError("Invalid signature")
        payload_bytes = base64.urlsafe_b64decode(encoded + b"=" * (-len(encoded) % 4))
        payload = json.loads(payload_bytes)
        if payload["exp"] <= int(datetime.now(timezone.utc).timestamp()):
            raise ValueError("Expired session")
        if payload["role"] not in {"superadmin", "admin"}:
            raise ValueError("Invalid role")
        if payload["role"] == "admin":
            with _admin_database() as connection:
                active_admin = connection.execute(
                    "SELECT 1 FROM admins WHERE username = ?",
                    (payload["sub"],),
                ).fetchone()
            if not active_admin:
                raise ValueError("Admin account no longer exists")
        return {"username": payload["sub"], "role": payload["role"]}
    except (ValueError, KeyError, TypeError, UnicodeDecodeError, json.JSONDecodeError):
        raise HTTPException(status_code=401, detail="Сессия истекла или недействительна.") from None


def _require_superadmin(request: Request) -> dict[str, str]:
    user = _current_user(request)
    if user["role"] != "superadmin":
        raise HTTPException(
            status_code=403,
            detail="Эта операция доступна только супер-администратору.",
        )
    return user


def _source_file(filename: str) -> Path:
    if not _is_safe_document_filename(filename):
        raise HTTPException(status_code=400, detail="Некорректное имя файла.")
    for directory in (UPLOAD_DIR, ORIGINAL_DIR):
        candidate = directory / filename
        if candidate.is_file():
            return candidate
    raise HTTPException(status_code=404, detail="Документ не найден.")


def _get_document_metadata(source: str) -> dict[str, str] | None:
    with _admin_database() as connection:
        row = connection.execute(
            "SELECT title, updated_at FROM document_metadata WHERE source = ?",
            (source,),
        ).fetchone()
    return dict(row) if row else None


def _set_document_metadata(source: str, title: str) -> None:
    updated_at = datetime.now().strftime("%d.%m.%Y")
    with _admin_database() as connection:
        connection.execute(
            """INSERT INTO document_metadata (source, title, updated_at)
            VALUES (?, ?, ?)
            ON CONFLICT(source) DO UPDATE SET
                title = excluded.title,
                updated_at = excluded.updated_at""",
            (source, title, updated_at),
        )


def _restore_document_metadata(source: str, metadata: dict[str, str] | None) -> None:
    with _admin_database() as connection:
        if metadata is None:
            connection.execute(
                "DELETE FROM document_metadata WHERE source = ?", (source,)
            )
        else:
            connection.execute(
                """INSERT INTO document_metadata (source, title, updated_at)
                VALUES (?, ?, ?)
                ON CONFLICT(source) DO UPDATE SET
                    title = excluded.title,
                    updated_at = excluded.updated_at""",
                (source, metadata["title"], metadata["updated_at"]),
            )


def _replace_file_contents(path: Path, content: bytes) -> None:
    temporary_path = path.with_name(f".{path.name}.{uuid.uuid4().hex}.tmp")
    try:
        temporary_path.write_bytes(content)
        temporary_path.replace(path)
    finally:
        temporary_path.unlink(missing_ok=True)


def _iter_table_paragraphs(table: Table):
    for row in table.rows:
        for cell in row.cells:
            yield from cell.paragraphs
            for nested_table in cell.tables:
                yield from _iter_table_paragraphs(nested_table)


def _iter_docx_paragraphs(document):
    yield from document.paragraphs
    for table in document.tables:
        yield from _iter_table_paragraphs(table)
    for section in document.sections:
        for part in (
            section.header,
            section.footer,
            section.first_page_header,
            section.first_page_footer,
            section.even_page_header,
            section.even_page_footer,
        ):
            yield from part.paragraphs
            for table in part.tables:
                yield from _iter_table_paragraphs(table)


def _iter_unique_docx_paragraphs(document):
    seen = set()
    for paragraph in _iter_docx_paragraphs(document):
        element_id = id(paragraph._p)
        if element_id not in seen:
            seen.add(element_id)
            yield paragraph


def _docx_placeholders(document) -> set[str]:
    pattern = re.compile(r"\{\{\s*([a-zA-Z][a-zA-Z0-9_]{0,49})\s*\}\}")
    return {
        match.group(1)
        for paragraph in _iter_docx_paragraphs(document)
        for match in pattern.finditer(paragraph.text)
    }


def _replace_docx_text_range(paragraph, start: int, end: int, replacement: str) -> None:
    runs = paragraph.runs
    run_starts = []
    offset = 0
    for run in runs:
        run_starts.append(offset)
        offset += len(run.text)
    start_index = next(
        index
        for index, run_start in enumerate(run_starts)
        if start < run_start + len(runs[index].text)
    )
    end_index = next(
        index
        for index, run_start in enumerate(run_starts)
        if end - 1 < run_start + len(runs[index].text)
    )
    start_offset = start - run_starts[start_index]
    end_offset = end - run_starts[end_index]
    if start_index == end_index:
        text = runs[start_index].text
        runs[start_index].text = text[:start_offset] + replacement + text[end_offset:]
        return

    runs[start_index].text = runs[start_index].text[:start_offset] + replacement
    for index in range(start_index + 1, end_index):
        runs[index].text = ""
    runs[end_index].text = runs[end_index].text[end_offset:]


def _apply_template_field_selections(fields_json: str, document) -> list[dict[str, str]]:
    try:
        fields = json.loads(fields_json)
    except json.JSONDecodeError as exc:
        raise HTTPException(status_code=400, detail="Проверьте список требуемых данных.") from exc
    if not isinstance(fields, list) or not fields or len(fields) > 50:
        raise HTTPException(
            status_code=400,
            detail="Выделите от 1 до 50 мест для заполнения в документе.",
        )

    paragraphs = list(_iter_unique_docx_paragraphs(document))
    validated_fields = []
    selections_by_paragraph: dict[int, list[tuple[int, int, str]]] = {}
    names = set()
    for field in fields:
        if not isinstance(field, dict):
            raise HTTPException(status_code=400, detail="Некорректное поле шаблона.")
        name = field.get("name")
        label = field.get("label")
        field_type = field.get("type")
        paragraph_index = field.get("paragraph_index")
        start = field.get("start")
        end = field.get("end")
        if (
            not isinstance(name, str)
            or not re.fullmatch(r"[a-zA-Z][a-zA-Z0-9_]{0,49}", name)
            or name in names
            or not isinstance(label, str)
            or not label.strip()
            or len(label) > 100
            or not isinstance(field_type, str)
            or field_type not in {"text", "date"}
            or not isinstance(paragraph_index, int)
            or isinstance(paragraph_index, bool)
            or not isinstance(start, int)
            or isinstance(start, bool)
            or not isinstance(end, int)
            or isinstance(end, bool)
            or paragraph_index < 0
            or paragraph_index >= len(paragraphs)
            or start < 0
            or end <= start
            or end > len(paragraphs[paragraph_index].text)
        ):
            raise HTTPException(
                status_code=400,
                detail="Проверьте выделенные фрагменты, названия полей и тип данных.",
            )
        selected_text = paragraphs[paragraph_index].text[start:end]
        if not selected_text.strip() or "{{" in selected_text or "}}" in selected_text:
            raise HTTPException(
                status_code=400,
                detail="Выберите обычный текст документа, не содержащий метки шаблона.",
            )
        names.add(name)
        selections_by_paragraph.setdefault(paragraph_index, []).append((start, end, name))
        validated_fields.append(
            {"name": name, "label": label.strip(), "type": field_type}
        )

    for paragraph_index, selections in selections_by_paragraph.items():
        ordered = sorted(selections)
        if any(current[0] < previous[1] for previous, current in zip(ordered, ordered[1:])):
            raise HTTPException(
                status_code=400,
                detail="Выбранные места для полей не должны пересекаться.",
            )
        for start, end, name in sorted(selections, reverse=True):
            _replace_docx_text_range(paragraphs[paragraph_index], start, end, f"{{{{{name}}}}}")
    return validated_fields


def _validate_template_fields(fields_json: str, document) -> list[dict[str, str]]:
    try:
        fields = json.loads(fields_json)
    except json.JSONDecodeError as exc:
        raise HTTPException(status_code=400, detail="Проверьте список требуемых данных.") from exc
    if not isinstance(fields, list) or len(fields) > 50:
        raise HTTPException(status_code=400, detail="Список данных должен содержать не более 50 полей.")

    validated_fields = []
    names = set()
    for field in fields:
        if not isinstance(field, dict):
            raise HTTPException(status_code=400, detail="Некорректное поле шаблона.")
        name = field.get("name")
        label = field.get("label")
        field_type = field.get("type")
        if (
            not isinstance(name, str)
            or not re.fullmatch(r"[a-zA-Z][a-zA-Z0-9_]{0,49}", name)
            or not isinstance(label, str)
            or not label.strip()
            or len(label) > 100
            or not isinstance(field_type, str)
            or field_type not in {"text", "date"}
        ):
            raise HTTPException(
                status_code=400,
                detail="Для каждого поля укажите название, метку и тип «текст» или «дата».",
            )
        if name in names:
            raise HTTPException(status_code=400, detail=f"Поле «{name}» указано несколько раз.")
        names.add(name)
        validated_fields.append(
            {"name": name, "label": label.strip(), "type": field_type}
        )

    placeholders = _docx_placeholders(document)
    if placeholders != names:
        missing_fields = sorted(placeholders - names)
        unused_fields = sorted(names - placeholders)
        details = []
        if missing_fields:
            details.append(f"добавьте поля для меток: {', '.join(missing_fields)}")
        if unused_fields:
            details.append(f"добавьте в DOCX метки: {', '.join('{{' + name + '}}' for name in unused_fields)}")
        raise HTTPException(
            status_code=400,
            detail="Не совпадают поля формы и метки в DOCX: " + "; ".join(details) + ".",
        )
    return validated_fields


def _replace_docx_placeholders(paragraph, values: dict[str, str]) -> None:
    runs = paragraph.runs
    if not runs:
        return
    combined_text = "".join(run.text for run in runs)
    matches = list(re.finditer(r"\{\{\s*([a-zA-Z][a-zA-Z0-9_]{0,49})\s*\}\}", combined_text))
    for match in reversed(matches):
        start, end = match.span()
        replacement = values[match.group(1)]
        run_starts = []
        offset = 0
        for run in runs:
            run_starts.append(offset)
            offset += len(run.text)

        start_index = next(
            index for index, run_start in enumerate(run_starts)
            if start < run_start + len(runs[index].text)
        )
        end_position = end - 1
        end_index = next(
            index for index, run_start in enumerate(run_starts)
            if end_position < run_start + len(runs[index].text)
        )
        start_offset = start - run_starts[start_index]
        end_offset = end - run_starts[end_index]
        if start_index == end_index:
            text = runs[start_index].text
            runs[start_index].text = text[:start_offset] + replacement + text[end_offset:]
            continue

        runs[start_index].text = runs[start_index].text[:start_offset] + replacement
        for index in range(start_index + 1, end_index):
            runs[index].text = ""
        runs[end_index].text = runs[end_index].text[end_offset:]


def _template_metadata(row: sqlite3.Row) -> dict:
    return {
        "id": row["template_id"],
        "title": row["title"],
        "description": row["description"],
        "category": row["category"],
        "fields": json.loads(row["fields_json"]),
    }


def _is_safe_document_filename(filename: str) -> bool:
    path = Path(filename)
    if (
        not filename
        or path.name != filename
        or path.suffix.lower() not in ALLOWED_EXTENSIONS
        or filename.rstrip(" .") != filename
        or re.fullmatch(r"(CON|PRN|AUX|NUL|COM[1-9]|LPT[1-9])", path.stem, flags=re.IGNORECASE)
    ):
        return False
    return all(
        char in "_ .()-" or unicodedata.category(char)[0] in {"L", "N", "M"}
        for char in filename
    )


@app.post("/api/chat")
def chat(request: ChatRequest, http_request: Request):
    _enforce_rate_limit(http_request, "chat", limit=20, window_seconds=60)
    sensitive_fields = _find_sensitive_data(
        "\n".join(
            [request.question, *(item.content for item in request.history)]
        )
    )
    if sensitive_fields:
        detected = ", ".join(sensitive_fields)
        raise HTTPException(
            status_code=400,
            detail=(
                f"Обнаружены возможные персональные данные ({detected}). "
                "Запрос не передан AI-сервису. Удалите эти данные и повторите попытку; "
                "если они уже есть в истории диалога, начните новый диалог."
            ),
        )
    try:
        return ask_bot(
            request.question,
            [{"role": item.role, "content": item.content} for item in request.history],
        )
    except Exception as exc:
        logger.exception("Chat request failed")
        raise HTTPException(
            status_code=503 if _embedding_quota_error_detail(exc) else 502,
            detail=_chat_error_detail(exc),
        ) from exc


@app.post("/api/auth/login")
def login(payload: LoginRequest, request: Request):
    _enforce_rate_limit(request, "login", limit=5, window_seconds=15 * 60)
    superadmin_username = os.getenv("SUPERADMIN_USERNAME", "").strip()
    superadmin_password = os.getenv("SUPERADMIN_PASSWORD", "")
    if not superadmin_username or len(superadmin_password) < 14:
        raise HTTPException(
            status_code=503,
            detail="Задайте SUPERADMIN_USERNAME и пароль SUPERADMIN_PASSWORD длиной не менее 14 символов в .env.",
        )
    _session_secret()

    if (
        hmac.compare_digest(payload.username.casefold(), superadmin_username.casefold())
        and hmac.compare_digest(payload.password, superadmin_password)
    ):
        user = {"username": superadmin_username, "role": "superadmin"}
    else:
        with _admin_database() as connection:
            admin = connection.execute(
                "SELECT username, password_hash FROM admins WHERE username = ?",
                (payload.username,),
            ).fetchone()
        if not admin or not _verify_admin_password(payload.password, admin["password_hash"]):
            raise HTTPException(status_code=401, detail="Неверный логин или пароль.")
        user = {"username": admin["username"], "role": "admin"}

    response = JSONResponse({"authenticated": True, **user})
    response.set_cookie(
        SESSION_COOKIE,
        _issue_session(user["username"], user["role"]),
        max_age=SESSION_DURATION_SECONDS,
        httponly=True,
        secure=request.url.scheme == "https",
        samesite="strict",
        path="/",
    )
    return response


@app.get("/api/auth/me")
def auth_status(request: Request):
    try:
        return {"authenticated": True, **_current_user(request)}
    except HTTPException as exc:
        if exc.status_code == 401:
            return {"authenticated": False}
        raise


@app.post("/api/auth/logout")
def logout():
    response = JSONResponse({"authenticated": False})
    response.delete_cookie(SESSION_COOKIE, path="/", httponly=True, samesite="strict")
    return response


@app.get("/api/admins")
def list_admins(request: Request):
    _require_superadmin(request)
    with _admin_database() as connection:
        admins = connection.execute(
            "SELECT username, created_at, created_by FROM admins ORDER BY username COLLATE NOCASE"
        ).fetchall()
    return {"admins": [dict(admin) for admin in admins]}


@app.post("/api/admins", status_code=201)
def create_admin(payload: AdminCreateRequest, request: Request):
    superadmin = _require_superadmin(request)
    reserved_username = os.getenv("SUPERADMIN_USERNAME", "").strip()
    if payload.username.casefold() == reserved_username.casefold():
        raise HTTPException(status_code=409, detail="Это имя зарезервировано за супер-администратором.")

    try:
        with _admin_database() as connection:
            connection.execute(
                "INSERT INTO admins (username, password_hash, created_at, created_by) VALUES (?, ?, ?, ?)",
                (
                    payload.username,
                    _admin_password_hash(payload.password),
                    datetime.now(timezone.utc).isoformat(),
                    superadmin["username"],
                ),
            )
    except sqlite3.IntegrityError as exc:
        raise HTTPException(status_code=409, detail="Пользователь с таким логином уже существует.") from exc
    return {"username": payload.username, "role": "admin"}


@app.delete("/api/admins/{username}", status_code=204)
def delete_admin(username: str, request: Request):
    _require_superadmin(request)
    if username.casefold() == os.getenv("SUPERADMIN_USERNAME", "").strip().casefold():
        raise HTTPException(status_code=400, detail="Учётную запись супер-администратора нельзя удалить.")
    with _admin_database() as connection:
        result = connection.execute("DELETE FROM admins WHERE username = ?", (username,))
    if result.rowcount == 0:
        raise HTTPException(status_code=404, detail="Администратор не найден.")


@app.get("/api/documents")
def list_documents(response: Response):
    documents = []
    paths_by_source = {}
    for directory in (ORIGINAL_DIR, UPLOAD_DIR):
        if directory.exists():
            for path in directory.iterdir():
                if path.is_file() and path.suffix.lower() in ALLOWED_EXTENSIONS:
                    paths_by_source[path.name] = path
    with _admin_database() as connection:
        metadata_by_source = {
            row["source"]: dict(row)
            for row in connection.execute(
                "SELECT source, title, updated_at FROM document_metadata"
            )
        }
    for source, path in paths_by_source.items():
        metadata = metadata_by_source.get(source)
        documents.append({
            "title": metadata["title"] if metadata else path.name,
            "description": "Доступен для поиска в чате.",
            "date": (
                metadata["updated_at"]
                if metadata
                else datetime.fromtimestamp(path.stat().st_mtime).strftime("%d.%m.%Y")
            ),
            "source": path.name,
            "indexed": True,
            "uploaded": path.parent.resolve() == UPLOAD_DIR.resolve(),
        })
    documents.sort(key=lambda item: item["title"].lower())
    response.headers["Cache-Control"] = "no-store"
    result: dict[str, object] = {"documents": documents}
    if not documents:
        result["detail"] = (
            "Документы не найдены. Проверьте DOCUMENTS_ORIGINAL_DIR и "
            "DOCUMENTS_UPLOAD_DIR, затем убедитесь, что в папке исходных "
            "документов есть DOCX или PDF."
        )
    return result


@app.get("/api/documents/{filename}/download")
def download_document(filename: str):
    path = _source_file(filename)
    return FileResponse(path, filename=path.name)


@app.get("/api/documents/{filename}/view")
def view_document(filename: str):
    path = _source_file(filename)
    media_type = (
        "application/pdf"
        if path.suffix.lower() == ".pdf"
        else "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    )
    return FileResponse(
        path,
        media_type=media_type,
        filename=path.name,
        content_disposition_type="inline",
    )


@app.get("/api/templates")
def list_templates(response: Response):
    with _admin_database() as connection:
        rows = connection.execute(
            "SELECT template_id, title, description, category, fields_json "
            "FROM templates ORDER BY title COLLATE NOCASE"
        ).fetchall()
    response.headers["Cache-Control"] = "no-store"
    return {"templates": [_template_metadata(row) for row in rows]}


@app.post("/api/templates/inspect")
def inspect_template(request: Request, file: UploadFile = File(...)):
    _current_user(request)
    filename = Path((file.filename or "").replace("\\", "/")).name
    if Path(filename).suffix.lower() != ".docx":
        raise HTTPException(status_code=400, detail="Шаблон должен быть файлом DOCX.")
    content = file.file.read(MAX_UPLOAD_SIZE + 1)
    if len(content) > MAX_UPLOAD_SIZE:
        raise HTTPException(
            status_code=413,
            detail=f"Размер файла превышает {MAX_UPLOAD_SIZE // (1024 * 1024)} МБ.",
        )
    if not content:
        raise HTTPException(status_code=400, detail="Файл шаблона пуст.")
    try:
        document = Document(BytesIO(content))
        paragraphs = [
            {"index": index, "text": paragraph.text}
            for index, paragraph in enumerate(_iter_unique_docx_paragraphs(document))
            if paragraph.text.strip()
        ]
    except Exception as exc:
        raise HTTPException(
            status_code=400,
            detail="Не удалось прочитать DOCX. Проверьте, что шаблон не повреждён.",
        ) from exc
    if not paragraphs:
        raise HTTPException(status_code=400, detail="В DOCX не найден текст для разметки.")
    if len(paragraphs) > 2000 or sum(len(item["text"]) for item in paragraphs) > 1_000_000:
        raise HTTPException(
            status_code=413,
            detail="Слишком большой документ для предпросмотра. Загрузите более короткий шаблон.",
        )
    return {"paragraphs": paragraphs}


@app.post("/api/templates", status_code=201)
def upload_template(
    request: Request,
    file: UploadFile = File(...),
    title: str = Form(...),
    description: str = Form(""),
    category: str = Form("Общие"),
    fields: str = Form(...),
):
    _current_user(request)
    filename = Path((file.filename or "").replace("\\", "/")).name
    if Path(filename).suffix.lower() != ".docx":
        raise HTTPException(status_code=400, detail="Шаблон должен быть файлом DOCX.")
    if (
        not title.strip()
        or len(title.strip()) > 100
        or len(description) > 500
        or not category.strip()
        or len(category.strip()) > 60
    ):
        raise HTTPException(
            status_code=400,
            detail="Название должно содержать до 100 символов, описание — до 500, категория — до 60 символов.",
        )
    content = file.file.read(MAX_UPLOAD_SIZE + 1)
    if len(content) > MAX_UPLOAD_SIZE:
        raise HTTPException(
            status_code=413,
            detail=f"Размер файла превышает {MAX_UPLOAD_SIZE // (1024 * 1024)} МБ.",
        )
    if not content:
        raise HTTPException(status_code=400, detail="Файл шаблона пуст.")
    try:
        document = Document(BytesIO(content))
        try:
            submitted_fields = json.loads(fields)
        except json.JSONDecodeError as exc:
            raise HTTPException(status_code=400, detail="Проверьте список требуемых данных.") from exc
        is_visual_setup = isinstance(submitted_fields, list) and any(
            isinstance(item, dict) and "paragraph_index" in item
            for item in submitted_fields
        )
        if is_visual_setup:
            validated_fields = _apply_template_field_selections(fields, document)
            fields = json.dumps(validated_fields, ensure_ascii=False)
            prepared_document = BytesIO()
            document.save(prepared_document)
            content = prepared_document.getvalue()
        else:
            validated_fields = _validate_template_fields(fields, document)
    except HTTPException:
        raise
    except Exception as exc:
        raise HTTPException(
            status_code=400,
            detail="Не удалось прочитать DOCX. Проверьте, что шаблон не повреждён.",
        ) from exc

    template_id = str(uuid.uuid4())
    TEMPLATE_DIR.mkdir(parents=True, exist_ok=True)
    template_path = TEMPLATE_DIR / f"{template_id}.docx"
    try:
        template_path.write_bytes(content)
        with _admin_database() as connection:
            connection.execute(
                "INSERT INTO templates "
                "(template_id, title, description, category, fields_json, created_at) "
                "VALUES (?, ?, ?, ?, ?, ?)",
                (
                    template_id,
                    title.strip(),
                    description.strip(),
                    category.strip(),
                    json.dumps(validated_fields, ensure_ascii=False),
                    datetime.now(timezone.utc).isoformat(),
                ),
            )
    except Exception:
        template_path.unlink(missing_ok=True)
        raise
    return {
        "template": {
            "id": template_id,
            "title": title.strip(),
            "description": description.strip(),
            "category": category.strip(),
            "fields": validated_fields,
        }
    }


@app.delete("/api/templates/{template_id}", status_code=204)
def delete_template(template_id: str, request: Request):
    _current_user(request)
    with _admin_database() as connection:
        row = connection.execute(
            "SELECT template_id FROM templates WHERE template_id = ?",
            (template_id,),
        ).fetchone()
        if not row:
            raise HTTPException(status_code=404, detail="Шаблон не найден.")
        (TEMPLATE_DIR / f"{template_id}.docx").unlink(missing_ok=True)
        connection.execute("DELETE FROM templates WHERE template_id = ?", (template_id,))


@app.post("/api/templates/{template_id}/generate")
def generate_template(template_id: str, payload: TemplateGenerateRequest):
    with _admin_database() as connection:
        row = connection.execute(
            "SELECT template_id, title, description, category, fields_json "
            "FROM templates WHERE template_id = ?",
            (template_id,),
        ).fetchone()
    if not row:
        raise HTTPException(status_code=404, detail="Шаблон не найден.")

    fields = json.loads(row["fields_json"])
    expected_names = {field["name"] for field in fields}
    if set(payload.values) != expected_names:
        raise HTTPException(
            status_code=400,
            detail="Переданные данные не совпадают с обязательными полями шаблона.",
        )
    values = {}
    for field in fields:
        value = payload.values[field["name"]].strip()
        if not value or len(value) > 4000:
            raise HTTPException(
                status_code=400,
                detail=f"Заполните поле «{field['label']}» (не более 4000 символов).",
            )
        if field["type"] == "date":
            try:
                value = datetime.strptime(value, "%Y-%m-%d").strftime("%d.%m.%Y")
            except ValueError as exc:
                raise HTTPException(
                    status_code=400,
                    detail=f"Укажите корректную дату в поле «{field['label']}».",
                ) from exc
        values[field["name"]] = value

    template_path = TEMPLATE_DIR / f"{template_id}.docx"
    if not template_path.is_file():
        raise HTTPException(status_code=404, detail="Файл шаблона не найден на сервере.")
    try:
        document = Document(template_path)
        for paragraph in _iter_docx_paragraphs(document):
            _replace_docx_placeholders(paragraph, values)
        if _docx_placeholders(document):
            raise ValueError("В документе остались незаполненные поля.")
        output = BytesIO()
        document.save(output)
    except Exception as exc:
        logger.exception("Could not generate template %s", template_id)
        raise HTTPException(
            status_code=500,
            detail="Не удалось сформировать DOCX из шаблона.",
        ) from exc

    safe_title = re.sub(r'[\\/:*?"<>|]+', "-", row["title"]).strip(" .-") or "document"
    download_name = f"{safe_title}.docx"
    headers = {
        "Content-Disposition": f"attachment; filename*=UTF-8''{quote(download_name)}"
    }
    return StreamingResponse(
        BytesIO(output.getvalue()),
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        headers=headers,
    )


@app.post("/api/documents/upload")
def upload_document(
    request: Request,
    file: UploadFile = File(...),
    title: str = Form(""),
):
    _current_user(request)
    filename = Path((file.filename or "").replace("\\", "/")).name
    suffix = Path(filename).suffix.lower()
    if not _is_safe_document_filename(filename):
        raise HTTPException(status_code=400, detail="Поддерживаются только файлы DOCX и PDF.")

    content = file.file.read(MAX_UPLOAD_SIZE + 1)
    if len(content) > MAX_UPLOAD_SIZE:
        raise HTTPException(
            status_code=413,
            detail=f"Размер файла превышает {MAX_UPLOAD_SIZE // (1024 * 1024)} МБ.",
        )
    if not content:
        raise HTTPException(status_code=400, detail="Файл пуст.")

    try:
        text = read_docx_bytes(content) if suffix == ".docx" else read_pdf_bytes(content)
    except Exception as exc:
        raise HTTPException(
            status_code=400,
            detail="Не удалось прочитать файл. Проверьте, что он не повреждён и соответствует своему формату.",
        ) from exc

    sensitive_fields = _find_sensitive_data(text)
    if sensitive_fields:
        detected = ", ".join(sensitive_fields)
        raise HTTPException(
            status_code=400,
            detail=(
                f"Индексация остановлена: в документе обнаружены возможные персональные "
                f"данные ({detected}). Файл не сохранён и не отправлен сервису эмбеддингов. "
                "Проверьте документ вручную и загрузите версию без персональных данных. "
                "Автоматическая проверка может ошибаться и не обнаруживает все виды данных."
            ),
        )

    chunks = create_splitter().split_text(text)
    if not chunks:
        raise HTTPException(status_code=400, detail="В файле не найден текст для индексации.")

    UPLOAD_DIR.mkdir(parents=True, exist_ok=True)
    target_path = UPLOAD_DIR / filename
    original_path = ORIGINAL_DIR / filename
    if target_path.exists() or original_path.is_file():
        raise HTTPException(
            status_code=409,
            detail="Документ с таким именем уже есть в базе. Обновлять существующие документы может только супер-администратор.",
        )

    try:
        client = QdrantClient(url=QDRANT_URL, trust_env=False)
        collection_exists = ensure_active_collection(client)
    except Exception as exc:
        logger.exception("Could not connect to Qdrant while uploading %s", filename)
        raise HTTPException(
            status_code=503,
            detail="Не удалось подключиться к Qdrant. Проверьте, что сервис запущен.",
        ) from exc
    if not collection_exists:
        raise HTTPException(
            status_code=409,
            detail="База знаний ещё не создана. Сначала выполните индексацию исходных документов.",
        )

    document_id = str(uuid.uuid4())
    clean_title = re.sub(r"[\r\n\t]+", " ", title).strip() or Path(filename).stem
    if len(clean_title) > 200:
        raise HTTPException(status_code=400, detail="Название документа не должно превышать 200 символов.")
    version = extract_version(text)
    metadatas = [{
        "source": filename,
        "title": clean_title,
        "version": version,
        "reference": extract_reference(chunk),
        "chunk": index,
        "file_path": str(target_path),
        "document_id": document_id,
    } for index, chunk in enumerate(chunks)]
    ids = [str(uuid.uuid4()) for _ in chunks]

    with DOCUMENT_INDEX_LOCK:
        if target_path.exists() or original_path.is_file():
            raise HTTPException(
                status_code=409,
                detail="Документ с таким именем уже есть в базе. Обновлять существующие документы может только супер-администратор.",
            )
        try:
            target_path.write_bytes(content)
            _set_document_metadata(filename, clean_title)
            vector_store = QdrantVectorStore.from_existing_collection(
                embedding=get_embeddings(),
                collection_name=ACTIVE_COLLECTION_NAME,
                url=QDRANT_URL,
                trust_env=False,
            )
            vector_store.add_texts(texts=chunks, metadatas=metadatas, ids=ids)
        except Exception as exc:
            target_path.unlink(missing_ok=True)
            _restore_document_metadata(filename, None)
            logger.exception("Document indexing failed for %s", filename)
            raise HTTPException(
                status_code=503 if _embedding_quota_error_detail(exc) else 502,
                detail=_embedding_quota_error_detail(exc)
                or "Не удалось проиндексировать документ. Проверьте доступность Qdrant и API эмбеддингов.",
            ) from exc

    return {
        "document": {
            "title": clean_title,
            "description": "Загружен и проиндексирован для поиска в чате.",
            "date": datetime.now().strftime("%d.%m.%Y"),
            "source": filename,
            "indexed": True,
            "uploaded": True,
        },
        "chunks": len(chunks),
        "message": f"Документ проиндексирован: {len(chunks)} фрагментов.",
    }


@app.put("/api/documents/{filename}")
def update_document(
    filename: str,
    request: Request,
    file: UploadFile = File(...),
    title: str = Form(""),
):
    _require_superadmin(request)
    target_path = _source_file(filename)
    if target_path.is_symlink():
        raise HTTPException(status_code=400, detail="Нельзя обновлять документ, являющийся символической ссылкой.")
    if target_path.suffix.lower() != ".docx":
        raise HTTPException(
            status_code=400,
            detail="Обновлять можно только документы DOCX.",
        )

    uploaded_filename = Path((file.filename or "").replace("\\", "/")).name
    if Path(uploaded_filename).suffix.lower() != ".docx":
        raise HTTPException(status_code=400, detail="Для обновления выберите файл DOCX.")
    content = file.file.read(MAX_UPLOAD_SIZE + 1)
    if len(content) > MAX_UPLOAD_SIZE:
        raise HTTPException(
            status_code=413,
            detail=f"Размер файла превышает {MAX_UPLOAD_SIZE // (1024 * 1024)} МБ.",
        )
    if not content:
        raise HTTPException(status_code=400, detail="Файл пуст.")

    try:
        text = read_docx_bytes(content)
    except Exception as exc:
        raise HTTPException(
            status_code=400,
            detail="Не удалось прочитать DOCX. Проверьте, что файл не повреждён.",
        ) from exc

    sensitive_fields = _find_sensitive_data(text)
    if sensitive_fields:
        detected = ", ".join(sensitive_fields)
        raise HTTPException(
            status_code=400,
            detail=(
                f"Обновление остановлено: в документе обнаружены возможные персональные "
                f"данные ({detected}). Файл не сохранён и не отправлен сервису эмбеддингов. "
                "Проверьте документ вручную и загрузите версию без персональных данных."
            ),
        )
    if not create_splitter().split_text(text):
        raise HTTPException(status_code=400, detail="В файле не найден текст для индексации.")

    requested_title = re.sub(r"[\r\n\t]+", " ", title).strip()
    if len(requested_title) > 200:
        raise HTTPException(
            status_code=400,
            detail="Название документа не должно превышать 200 символов.",
        )

    with DOCUMENT_INDEX_LOCK:
        source_path = _source_file(filename)
        if source_path.is_symlink():
            raise HTTPException(
                status_code=400,
                detail="Нельзя обновлять документ, являющийся символической ссылкой.",
            )
        if source_path.suffix.lower() != ".docx":
            raise HTTPException(
                status_code=400,
                detail="Обновлять можно только документы DOCX.",
            )
        create_override = (
            source_path.parent.resolve() == ORIGINAL_DIR.resolve()
            and ORIGINAL_DIR.resolve() != UPLOAD_DIR.resolve()
        )
        target_path = UPLOAD_DIR / filename if create_override else source_path
        previous_content = source_path.read_bytes()
        previous_stat = source_path.stat()
        previous_metadata = _get_document_metadata(filename)
        clean_title = requested_title or (
            previous_metadata["title"] if previous_metadata else source_path.name
        )
        try:
            if create_override:
                UPLOAD_DIR.mkdir(parents=True, exist_ok=True)
            _replace_file_contents(target_path, content)
            _set_document_metadata(filename, clean_title)
            reindex_documents((ORIGINAL_DIR, UPLOAD_DIR), ADMIN_DATABASE)
        except Exception as exc:
            try:
                if create_override:
                    target_path.unlink(missing_ok=True)
                else:
                    _replace_file_contents(target_path, previous_content)
                    os.utime(
                        target_path,
                        ns=(previous_stat.st_atime_ns, previous_stat.st_mtime_ns),
                    )
                _restore_document_metadata(filename, previous_metadata)
            except Exception as rollback_error:
                logger.exception(
                    "Failed to restore document %s after update error",
                    filename,
                )
                raise HTTPException(
                    status_code=500,
                    detail="Обновление не завершилось, и автоматический откат не удался. Проверьте файл на сервере.",
                ) from rollback_error
            logger.exception("Document update failed for %s; previous version restored", filename)
            raise HTTPException(
                status_code=502,
                detail=(
                    "Не удалось обновить индекс документов. Исходный файл восстановлен; "
                    "проверьте доступность Qdrant и API эмбеддингов."
                ),
            ) from exc

    metadata = _get_document_metadata(filename)
    return {
        "document": {
            "title": metadata["title"] if metadata else target_path.name,
            "description": "Документ обновлён и переиндексирован для поиска в чате.",
            "date": (
                metadata["updated_at"]
                if metadata
                else datetime.fromtimestamp(target_path.stat().st_mtime).strftime("%d.%m.%Y")
            ),
            "source": filename,
            "indexed": True,
            "uploaded": True,
        },
        "message": "Документ обновлён и переиндексирован.",
    }


@app.delete("/api/documents/{filename}", status_code=204)
def delete_uploaded_document(filename: str, request: Request):
    _require_superadmin(request)
    with DOCUMENT_INDEX_LOCK:
        target_path = _source_file(filename)
        if target_path.is_symlink():
            raise HTTPException(
                status_code=400,
                detail="Нельзя удалить документ, являющийся символической ссылкой.",
            )
        if target_path.parent.resolve() != UPLOAD_DIR.resolve():
            raise HTTPException(
                status_code=400,
                detail="Удалять через интерфейс можно только загруженные документы.",
            )
        previous_content = target_path.read_bytes()
        previous_stat = target_path.stat()
        previous_metadata = _get_document_metadata(filename)
        try:
            target_path.unlink()
            _restore_document_metadata(filename, None)
            reindex_documents((ORIGINAL_DIR, UPLOAD_DIR), ADMIN_DATABASE)
        except Exception as exc:
            try:
                if not target_path.exists():
                    target_path.parent.mkdir(parents=True, exist_ok=True)
                    target_path.write_bytes(previous_content)
                    os.utime(
                        target_path,
                        ns=(previous_stat.st_atime_ns, previous_stat.st_mtime_ns),
                    )
                _restore_document_metadata(filename, previous_metadata)
            except Exception as rollback_error:
                logger.exception(
                    "Failed to restore deleted document %s after reindex error",
                    filename,
                )
                raise HTTPException(
                    status_code=500,
                    detail="Не удалось перестроить индекс, и автоматическое восстановление файла не удалось.",
                ) from rollback_error
            logger.exception(
                "Document deletion failed for %s; previous version restored",
                filename,
            )
            raise HTTPException(
                status_code=502,
                detail=(
                    "Не удалось удалить документ из индекса. Файл восстановлен; "
                    "проверьте доступность Qdrant и API эмбеддингов."
                ),
            ) from exc


@app.post("/api/documents/check")
def check_document(
    file: UploadFile = File(...),
):
    filename = Path((file.filename or "").replace("\\", "/")).name
    if Path(filename).suffix.lower() != ".docx":
        raise HTTPException(status_code=400, detail="Для проверки выберите файл DOCX.")
    content = file.file.read(MAX_UPLOAD_SIZE + 1)
    if len(content) > MAX_UPLOAD_SIZE:
        raise HTTPException(
            status_code=413,
            detail=f"Размер файла превышает {MAX_UPLOAD_SIZE // (1024 * 1024)} МБ.",
        )
    try:
        text = read_docx_bytes(content)
    except Exception as exc:
        raise HTTPException(
            status_code=400,
            detail="Не удалось прочитать DOCX. Проверьте, что файл не повреждён.",
        ) from exc
    if not text.strip():
        raise HTTPException(status_code=400, detail="В документе не найден текст для проверки.")
    sensitive_fields = _find_sensitive_data(text)
    if sensitive_fields:
        detected = ", ".join(sensitive_fields)
        raise HTTPException(
            status_code=400,
            detail=(
                f"Проверка остановлена: обнаружены возможные персональные данные "
                f"({detected}). Документ не передан внешнему AI-сервису. Удалите эти "
                "данные или проверьте файл вручную. Автоматическая проверка может "
                "ошибаться и не обнаруживает все виды данных."
            ),
        )
    try:
        return review_draft(text)
    except Exception as exc:
        logger.exception("Document check failed for %s", filename)
        raise HTTPException(
            status_code=503 if _embedding_quota_error_detail(exc) else 502,
            detail=_chat_error_detail(exc),
        ) from exc


@app.get("/")
def index():
    return FileResponse(APP_DIR / "index.html", headers={"Cache-Control": "no-store"})


@app.get("/styles.css")
def styles():
    return FileResponse(
        APP_DIR / "styles.css",
        media_type="text/css",
        headers={"Cache-Control": "no-store"},
    )


@app.get("/script.js")
def script():
    return FileResponse(
        APP_DIR / "script.js",
        media_type="application/javascript",
        headers={"Cache-Control": "no-store"},
    )
